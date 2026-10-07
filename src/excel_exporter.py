import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def export_invoices_to_excel(invoices, output_path):
    dataframe = pd.DataFrame(invoices)

    dataframe.to_excel(
        output_path,
        index=False,
        sheet_name="Invoices"
    )

    workbook = load_workbook(output_path)
    worksheet = workbook["Invoices"]

    # Format header
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Format money columns
    money_columns = ["Subtotal", "Tax", "Total"]

    for column_name in money_columns:
        if column_name in dataframe.columns:
            column_number = dataframe.columns.get_loc(column_name) + 1

            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(
                    row=row,
                    column=column_number
                ).number_format = '#,##0.00'

    # Minimum widths for invoice columns
    minimum_widths = {
        "Invoice Number": 18,
        "Invoice Date": 20,
        "Vendor": 28,
        "Customer": 24,
        "GSTIN": 18,
        "Subtotal": 14,
        "Tax": 14,
        "Total": 14
    }

    # Adjust column widths
    for column in worksheet.columns:
        column_letter = get_column_letter(column[0].column)
        column_name = worksheet.cell(
            row=1,
            column=column[0].column
        ).value

        max_length = 0

        for cell in column:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        width = max(
            max_length + 3,
            minimum_widths.get(column_name, 12)
        )

        worksheet.column_dimensions[column_letter].width = width

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Enable filters
    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(output_path)